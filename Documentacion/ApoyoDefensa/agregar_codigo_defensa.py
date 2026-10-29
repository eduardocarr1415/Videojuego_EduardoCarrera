from pathlib import Path
from textwrap import dedent
from docx import Document
from docx.shared import Pt
import shutil

root=Path(r'C:\Users\soyel\Marihuanero')
path=root/'Documentacion/Defensa_Marihuanero.docx'
backup=path.with_name('Defensa_Marihuanero_antes_de_ampliar.docx')
if not backup.exists(): shutil.copy2(path, backup)
d=Document(backup)

def p(text): d.add_paragraph(text)
def h(text): d.add_heading(text,2)
def snippet(file,start,end):
    src=(root/'Assets/Scripts'/file).read_text(encoding='utf-8-sig')
    a=src.index(start); b=src.index(end,a)+len(end)
    # Añadimos la sangría original para poder quitarla sin alterar el código.
    a=src.rfind('\n',0,a)+1
    lines=[ln for ln in src[a:b].splitlines() if not ln.lstrip().startswith('//')]
    text=dedent('\n'.join(lines)).strip()
    para=d.add_paragraph(); para.paragraph_format.line_spacing=1.0
    para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(9)
    r=para.add_run(text);r.font.name='Consolas';r.font.size=Pt(9.5)
    return text

def page(n,name,where):
    d.add_page_break();d.add_heading(f'{n}  Código de {name}',1)
    p(f'Archivo {name}.cs dentro de Assets/Scripts. {where}')

# Evitamos acumular el añadido si se ejecuta otra vez este archivo.
for para in d.paragraphs:
    if para.text.startswith('Primero entiende el recorrido completo:'):
        para.text += ' Las páginas 18 a 25 incluyen código real de los ocho scripts y frases para explicarlo en la defensa.'
    if para.text.startswith('cuerpo.linearVelocity = new Vector2('):
        para.text='Vector2 velocidad = new Vector2(\n    direccion * rapidez,\n    cuerpo.linearVelocity.y\n);'
        for r in para.runs:r.font.name='Consolas';r.font.size=Pt(9)

page(18,'ControlJugador','Está en Jugador. Estos fragmentos son partes de FixedUpdate y OnTriggerEnter2D.')
h('Preparar el movimiento y aceptar el salto')
snippet('ControlJugador.cs','Vector2 velocidad = new Vector2(','saltoPendiente = false;')
p('Primero preparo dos números: el avance a los lados y la velocidad vertical que ya tenía. La pregunta if exige que se haya pedido el salto y que haya piso debajo. El signo && significa que deben cumplirse las dos cosas. Entonces cambio la velocidad vertical, pido el sonido y entrego la velocidad al cuerpo. Al final borro la petición de salto.')
h('Reconocer una hoja')
snippet('ControlJugador.cs','if (otro.CompareTag("Hoja"))','Destroy(otro.gameObject);\n        }')
p('Reconozco la etiqueta Hoja. Apago su caja para no contarla dos veces, aviso al administrador y borro esa copia. Destroy no elimina el molde de Project.')
h('Así lo explico al profesor')
p('“Este bloque mueve el cuerpo y solo le permite saltar cuando está apoyado. El otro bloque reconoce una hoja y le avisa al encargado del nivel para sumarla. El personaje reconoce el contacto; el administrador guarda la cuenta”.')

page(19,'ControlNivel','Está en AdministradorNivel. Aquí se cuentan las hojas y se decide el resultado.')
h('Sumar una hoja y cambiar el texto')
snippet('ControlNivel.cs','public void RecogerHoja()','textoHojas.text = "Hojas: " + hojasRecogidas + " / " + totalHojas;\n    }')
p('Si ya terminó, return sale del método. Si seguimos jugando, ++ suma una hoja. Después pide el sonido y actualiza el texto. Los signos + unen palabras y números para escribir la cuenta. El signo ?. pide ejecutar el sonido solo si esa conexión existe.')
h('Lo que ocurre al perder')
snippet('ControlNivel.cs','public void Morir()','Time.timeScale = 0f;\n    }')
p('La primera pregunta evita repetir la derrota. Después guardo que perdimos, pido el sonido, enciendo el panel y dejo quieto el juego. La victoria sigue una idea parecida, pero antes compara hojasRecogidas con totalHojas dentro de IntentarGanar.')
h('Así lo explico al profesor')
p('“ControlNivel es el encargado de las reglas. No lee el teclado: recibe avisos. Al recoger una hoja suma una y cambia el contador. Cuando perdemos, muestra el resultado una sola vez y pausa el movimiento”.')

page(20,'AnimarJugador','Está en Jugador y controla el dibujo de ImagenJugador. Estos bloques están en MostrarMovimiento y MostrarPoseSalto.')
h('Elegir entre caminar y estar quieto')
snippet('AnimarJugador.cs','int siguiente = Mathf.Abs(velocidad.x)','estadoActual = siguiente;\n                }')
p('Mathf.Abs mira cuánto se mueve hacia un lado sin importar si es izquierda o derecha. Si supera 0,05 elige Caminar; si no, Quieto. La parte con ? y : significa “si pasa esto elige lo primero y si no elige lo segundo”. Luego evita comenzar de nuevo una animación que ya estaba sonando.')
h('Mostrar una pose del salto')
snippet('AnimarJugador.cs','void MostrarPoseSalto(float fase)','estadoActual = Saltar;\n    }')
p('Pongo la velocidad del animador en cero para dejar fija una pose del salto. Play elige Saltar y la parte indicada por fase. Esto no congela el cuerpo: solo el dibujo. El personaje sigue subiendo o cayendo con su física.')
h('Así lo explico al profesor')
p('“Primero miro cómo se está moviendo el cuerpo y luego elijo la imagen. Si camina cambio a Caminar; si se detiene uso Quieto. En el salto elijo una pose que combine con su subida o caída. Este script no le da fuerza al personaje”.')

page(21,'CamaraSigue','Está en Main Camera. Los fragmentos pertenecen a LateUpdate.')
h('Acompañar el avance del personaje')
snippet('CamaraSigue.cs','float avance = jugador.position.x','transform.position += new Vector3(avance, 0f, 0f);')
p('Resto la posición de antes a la de ahora para saber cuánto avanzó. Guardo la posición nueva para la próxima revisión. Después muevo la cámara esa misma distancia a los lados. Como los otros dos valores son cero, este bloque no cambia ni su altura ni su profundidad.')
h('Corregir el encuadre con suavidad')
snippet('CamaraSigue.cs','Vector3 destino = new Vector3(','suavizado\n        );')
p('El destino queda dos unidades delante del personaje. Conservo la Y y la Z de la cámara. SmoothDamp acerca la cámara a ese destino de forma suave cuando hay una diferencia. suavizado indica cómo se hace esa corrección; velocidad guarda un dato que necesita ese cálculo entre revisiones.')
h('Así lo explico al profesor')
p('“La cámara acompaña lo que avanzó el jugador y deja ver un poco de camino por delante. No sube con cada salto porque conserva su altura. Eso ayuda a que el personaje no parezca retroceder por culpa del encuadre”.')
p('Para acercar la vista se cambió Size en el componente Camera, no en este bloque. El valor guardado es 4,2. Es un ajuste del Inspector que trabaja junto al seguimiento.')

page(22,'AjustarFondo','Se usa en los fondos con Sprite Renderer. Este fragmento está dentro de AjustarTamano.')
h('Calcular cuánto debe crecer la imagen')
snippet('AjustarFondo.cs','float altoPantalla = camara.orthographicSize','altoPantalla / tamanoImagen.y\n        ) * 1.002f;')
p('orthographicSize es la mitad de la altura que ve la cámara, por eso lo multiplico por dos. aspect dice la relación entre ancho y alto. Las divisiones comparan el espacio visible con el tamaño del dibujo. Mathf.Max elige el aumento mayor para que la imagen cubra toda la pantalla. El 1,002 deja un margen muy pequeño en los bordes.')
h('Aplicar el cambio solo cuando hace falta')
snippet('AjustarFondo.cs','if ((transform.localScale - nuevoTamano)','transform.localScale = nuevoTamano;')
p('nuevoTamano se preparó antes teniendo en cuenta el tamaño del objeto padre. Esta pregunta compara ese tamaño con el que ya tiene la imagen. Si son prácticamente iguales no vuelve a cambiarlo. No es necesario memorizar el número pequeño: sirve como margen para decidir si la diferencia importa.')
h('Así lo explico al profesor')
p('“Este bloque calcula qué tan grande debe ser el fondo para llenar la vista sin deformar el dibujo. Usa el mismo aumento para ancho y alto. Puede recortar un borde de la imagen, pero evita dejar un espacio vacío alrededor”.')
p('Si quiero demostrarlo, cambio la forma de la ventana Game y señalo que el fondo sigue cubriendo la vista. Esto modifica la apariencia; no cambia las cajas del piso ni el movimiento del personaje.')

page(23,'AnimarSelva','Está en la capa AtmosferaSelva del nivel. Da valores al material que pinta la niebla y las luces.')
h('Hacer que el ciclo vuelva a empezar')
snippet('AnimarSelva.cs','void LateUpdate()','ActualizarCapa();\n    }')
p('Application.isPlaying pregunta si estamos probando el juego. Time.deltaTime dice cuánto tiempo pasó desde la última imagen. Lo sumo al reloj. Mathf.Repeat hace que el reloj vuelva al principio cuando llega a la duración del ciclo. Luego actualizo la apariencia. Al detenerse el tiempo del juego, este reloj deja de avanzar.')
h('Enviar los valores al dibujo')
snippet('AnimarSelva.cs','propiedades.SetFloat(Fase,','propiedades.SetFloat(Particulas, brilloParticulas);')
p('Aquí envío por qué parte del ciclo vamos, la forma de la pantalla y cuánto se ven niebla, rayos y puntitos. Estas líneas preparan los valores. Más abajo se aplican a la capa con SetPropertyBlock. El shader los recibe y decide los colores transparentes que se dibujan.')
h('Así lo explico al profesor')
p('“No usé un GIF. Tengo un reloj que se repite y le dice al material cómo variar la niebla, la luz y los puntitos. La imagen de selva es el fondo y esta capa le da movimiento. Las cajas del suelo siguen siendo las mismas”.')
p('La duración inicial del ciclo es diez segundos. También puedo cambiar la cantidad de niebla desde una casilla, sin dibujar otro fondo ni cambiar el recorrido.')

page(24,'SonidosJuego','Está en AdministradorNivel. Usa un reproductor para los efectos y una conexión a la música.')
h('Pedir un efecto al recoger una hoja')
snippet('SonidosJuego.cs','public void RecogerHoja()','fuente.PlayOneShot(hoja, volumenHoja);\n    }')
p('El signo ! significa “no”. La pregunta pide que la partida no haya terminado y que exista el audio de hoja. PlayOneShot lo reproduce una vez con el volumen elegido. La grabación hoja y el reproductor fuente son cosas distintas: uno es el archivo y el otro lo hace sonar.')
h('Escuchar el resultado una sola vez')
snippet('SonidosJuego.cs','private void ReproducirResultado(AudioClip sonido)','if (sonido != null) fuente.PlayOneShot(sonido, volumenResultado);\n    }')
p('Primero evito repetir el final. Después marco que terminó, detengo la música si existe y corto los efectos anteriores. Finalmente reproduzco la grabación recibida. Perder le entrega derrota y Ganar le entrega victoria. El mismo bloque sirve para los dos resultados.')
h('Así lo explico al profesor')
p('“Cuando recojo una hoja se pide un efecto corto. Al ganar o perder uso otro bloque que corta la música y deja escuchar el resultado una sola vez. La muerte suena pronto porque el archivo de audio tiene recortado el silencio inicial”.')
p('No hay una espera escondida en este método. El orden de las instrucciones permite pedir el sonido antes de mostrar el panel y detener el movimiento.')

page(25,'ControlMenuPrincipal','Está en AdministradorMenu dentro de MenuPrincipal. Los botones llaman a estos métodos.')
h('Abrir el nivel al pulsar JUGAR')
snippet('ControlMenuPrincipal.cs','public void Jugar()','SceneManager.LoadScene(escenaJuego);\n    }')
p('Si ya se pidió abrir el nivel, return evita repetir la orden. Si todavía no, guardo cargando como verdadero, pongo el tiempo normal y cargo la escena. escenaJuego es una casilla que contiene SampleScene. El nombre debe coincidir y la escena debe estar en la lista de Build Profiles.')
h('Salir según dónde se está ejecutando')
snippet('ControlMenuPrincipal.cs','public void Salir()','#endif\n    }')
p('Dentro de Unity se usa la parte UNITY_EDITOR para detener Play. En el juego exportado se usa la otra parte para cerrar la aplicación. Las líneas que empiezan con # hacen que Unity elija el bloque correspondiente al preparar el programa.')
h('Así lo explico al profesor')
p('“El menú no contiene las plataformas. Solo presenta el juego y abre la escena del nivel. Jugar carga SampleScene y Salir cierra la prueba o el juego según dónde lo esté usando. El botón NIVELES está reservado para una mejora y todavía no tiene acción”.')
h('Cómo estudio estos fragmentos')
p('Leo primero el nombre del bloque. Después digo qué dato recibe, qué pregunta hace y qué cambia en el juego. Si me piden más detalle, abro el archivo en Assets/Scripts y uso los comentarios que ya están escritos encima de cada parte.')

d.save(path)
print(path)
print('Se añadieron 8 páginas con fragmentos tomados de los scripts actuales')
